import io
from pathlib import Path
import subprocess
import unittest
from unittest.mock import patch
import zipfile

from docx import Document
from pypdf import PdfReader
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from modules import APP_MODULES
from services.access_control import account_can_access_module
from services.json_tools import JsonInputError, format_json, minify_json, parse_json
from services.pdf_tools import (
    ConversionUnavailableError,
    PdfToolError,
    docx_to_pdf,
    merge_pdfs,
    parse_page_ranges,
    pdf_to_docx,
    split_pdf,
    validate_bytes,
)
from services.qr_tools import (
    QR_ERROR_LEVELS,
    generate_qr_png,
    validate_text_payload,
    wifi_payload,
)


def create_pdf(page_count: int, *, text: str | None = "Test page") -> bytes:
    output = io.BytesIO()
    document = canvas.Canvas(output, pagesize=letter)
    for page_number in range(page_count):
        if text is not None:
            document.drawString(72, 720, f"{text} {page_number + 1}")
        document.showPage()
    document.save()
    return output.getvalue()


def create_docx(text: str = "Editable document") -> bytes:
    output = io.BytesIO()
    document = Document()
    document.add_paragraph(text)
    document.save(output)
    return output.getvalue()


class JsonToolTests(unittest.TestCase):
    def test_formats_and_minifies_json(self):
        source = '{"name":"Ada","active":true}'
        self.assertEqual(
            format_json(source),
            '{\n  "name": "Ada",\n  "active": true\n}',
        )
        self.assertEqual(minify_json(source), source)
        self.assertEqual(parse_json(source)["name"], "Ada")

    def test_reports_invalid_json_location_and_empty_input(self):
        with self.assertRaises(JsonInputError) as error:
            format_json('{"name":}')
        self.assertEqual((error.exception.line, error.exception.column), (1, 9))
        with self.assertRaisesRegex(JsonInputError, "Escribe"):
            parse_json("  ")
        with self.assertRaisesRegex(JsonInputError, "2 MB"):
            parse_json("x" * (2 * 1024 * 1024 + 1))

    def test_rejects_non_standard_nan(self):
        with self.assertRaisesRegex(JsonInputError, "no está permitido"):
            parse_json('{"value": NaN}')


class QrToolTests(unittest.TestCase):
    def test_generates_png_for_url(self):
        content = validate_text_payload("https://example.com")
        png = generate_qr_png(
            content,
            error_correction=QR_ERROR_LEVELS["Alto (25 %)"],
            box_size=5,
        )
        self.assertTrue(png.startswith(b"\x89PNG\r\n\x1a\n"))

    def test_validates_text_and_escapes_wifi_payload(self):
        self.assertEqual(
            wifi_payload("Office;Wi-Fi", "s;ecret", "WPA/WPA2"),
            r"WIFI:T:WPA;S:Office\;Wi-Fi;P:s\;ecret;;",
        )
        self.assertEqual(
            wifi_payload("Guest", "", "Sin contraseña"),
            "WIFI:T:nopass;S:Guest;;",
        )
        with self.assertRaisesRegex(ValueError, "http://"):
            validate_text_payload("ftp://example.com")
        with self.assertRaisesRegex(ValueError, "contraseña"):
            wifi_payload("Private", "", "WPA/WPA2")


class PdfToolTests(unittest.TestCase):
    def test_merges_documents_in_requested_order_and_splits_ranges(self):
        first = create_pdf(2, text="First")
        second = create_pdf(1, text="Second")
        merged = merge_pdfs([("second.pdf", second), ("first.pdf", first)])
        self.assertEqual(len(PdfReader(io.BytesIO(merged)).pages), 3)
        extracted = "\n".join(
            page.extract_text() or "" for page in PdfReader(io.BytesIO(merged)).pages
        )
        self.assertLess(extracted.index("Second 1"), extracted.index("First 1"))

        selected = split_pdf(first, "2,1")
        self.assertEqual(len(PdfReader(io.BytesIO(selected)).pages), 2)
        self.assertEqual(parse_page_ranges("1-2,4", 4), [0, 1, 3])

    def test_rejects_invalid_ranges_and_damaged_pdfs(self):
        for ranges in ("", "3-1", "1,1", "0", "1-5"):
            with self.subTest(ranges=ranges), self.assertRaises(PdfToolError):
                split_pdf(create_pdf(3), ranges)
        with self.assertRaises(PdfToolError):
            merge_pdfs([("a.pdf", b"not a pdf"), ("b.pdf", create_pdf(1))])
        with self.assertRaisesRegex(PdfToolError, "No repitas"):
            parse_page_ranges("1,1", 2)
        with self.assertRaises(PdfToolError):
            parse_page_ranges("1-999999999999999999999", 2)

    def test_pdf_to_docx_extracts_text_and_refuses_scanned_empty_pdf(self):
        result = pdf_to_docx(create_pdf(1, text="Editable text"))
        validate_bytes(result, kind="docx")
        converted = Document(io.BytesIO(result))
        self.assertIn("Editable text 1", converted.paragraphs[0].text)
        with self.assertRaisesRegex(PdfToolError, "OCR"):
            pdf_to_docx(create_pdf(1, text=None))

    def test_docx_validation_and_missing_libreoffice_message(self):
        source = create_docx()
        validate_bytes(source, kind="docx")
        with patch("services.pdf_tools.shutil.which", return_value=None):
            with self.assertRaises(ConversionUnavailableError):
                docx_to_pdf(source)
        with self.assertRaises(PdfToolError):
            validate_bytes(b"not a docx", kind="docx")

    def test_word_to_pdf_uses_and_cleans_a_private_temporary_directory(self):
        source_docx = create_docx()
        expected_pdf = create_pdf(1)
        temporary_directories = []

        def fake_run(command, **kwargs):
            output_dir = Path(command[command.index("--outdir") + 1])
            temporary_directories.append(output_dir)
            (output_dir / "documento.pdf").write_bytes(expected_pdf)
            return subprocess.CompletedProcess(command, 0, b"", b"")

        with (
            patch("services.pdf_tools.shutil.which", return_value="/usr/bin/libreoffice"),
            patch("services.pdf_tools.subprocess.run", side_effect=fake_run),
        ):
            converted = docx_to_pdf(source_docx)

        self.assertTrue(converted.startswith(b"%PDF-"))
        self.assertEqual(len(temporary_directories), 1)
        self.assertFalse(temporary_directories[0].exists())


class PermissionAndCatalogTests(unittest.TestCase):
    def test_checks_each_module_permission_and_enabled_state(self):
        user = {
            "active": True,
            "role": "USER",
            "module_ids": ["qr_generator"],
        }
        self.assertTrue(
            account_can_access_module(user, "qr_generator", module_enabled=True)
        )
        self.assertFalse(
            account_can_access_module(user, "pdf_toolbox", module_enabled=True)
        )
        self.assertFalse(
            account_can_access_module(user, "qr_generator", module_enabled=False)
        )
        self.assertFalse(
            account_can_access_module(
                user,
                "qr_generator",
                module_enabled=True,
                admin_only=True,
            )
        )
        admin = {**user, "role": "SUPER_ADMIN", "module_ids": []}
        self.assertTrue(
            account_can_access_module(admin, "pdf_toolbox", module_enabled=True)
        )
        self.assertFalse(
            account_can_access_module(
                {**admin, "active": False}, "pdf_toolbox", module_enabled=True
            )
        )

    def test_all_apps_are_registered_and_have_page_files(self):
        ids = {module["id"] for module in APP_MODULES}
        self.assertTrue(
            {"gestion", "reloj_pdf", "qr_generator", "json_formatter", "pdf_toolbox"}
            <= ids
        )
        for module in APP_MODULES:
            for page in module["pages"]:
                self.assertTrue(Path(page["path"]).is_file(), page["path"])

    def test_direct_page_entry_rechecks_server_side_permissions(self):
        from streamlit.testing.v1 import AppTest

        def fake_supabase(module_ids):
            class Query:
                def __init__(self, table):
                    self.table_name = table
                    self.filters = {}

                def select(self, _columns):
                    return self

                def eq(self, column, value):
                    self.filters[column] = value
                    return self

                def limit(self, _count):
                    return self

                def execute(self):
                    if self.table_name == "app_access_users":
                        rows = [
                            {
                                "email": self.filters.get("email"),
                                "role": "USER",
                                "module_ids": module_ids,
                                "active": True,
                            }
                        ]
                    else:
                        rows = [{"id": "pdf_toolbox", "enabled": True}]
                    return type("Response", (), {"data": rows})()

            class Supabase:
                def table(self, name):
                    return Query(name)

            return Supabase()

        def protected_page():
            import streamlit as st

            from services.access_control import require_module_access

            require_module_access("pdf_toolbox")
            st.title("PDF Toolbox autorizado")

        app = AppTest.from_function(protected_page)
        app.session_state["authenticated_email"] = "user@example.com"
        with patch(
            "services.access_control.init_connection",
            return_value=fake_supabase(["qr_generator"]),
        ):
            app.run()
        self.assertTrue(app.error)
        self.assertFalse(any("PDF Toolbox autorizado" in item.value for item in app.title))

        authorized = AppTest.from_function(protected_page)
        authorized.session_state["authenticated_email"] = "user@example.com"
        with patch(
            "services.access_control.init_connection",
            return_value=fake_supabase(["pdf_toolbox"]),
        ):
            authorized.run()
        self.assertFalse(authorized.exception)
        self.assertTrue(
            any("PDF Toolbox autorizado" in item.value for item in authorized.title)
        )

    def test_existing_hub_login_screen_renders(self):
        from streamlit.testing.v1 import AppTest

        class EmptyQuery:
            def select(self, _columns):
                return self

            def limit(self, _count):
                return self

            def execute(self):
                return type("Response", (), {"data": []})()

        class EmptySupabase:
            def table(self, _name):
                return EmptyQuery()

        root = Path(__file__).resolve().parents[1]
        with patch("services.db.init_connection", return_value=EmptySupabase()):
            app = AppTest.from_file(str(root / "app.py")).run()
        self.assertFalse(app.exception)
        self.assertTrue(any("Hub de aplicaciones" in item.value for item in app.title))
        self.assertEqual(len(app.text_input), 2)

    def test_hub_only_lists_modules_assigned_to_current_user(self):
        from streamlit.testing.v1 import AppTest

        app_ids = {
            module["id"]: module["title"]
            for module in APP_MODULES
        }
        account = {
            "email": "user@example.com",
            "display_name": "Test user",
            "role": "USER",
            "module_ids": ["qr_generator", "json_formatter"],
            "active": True,
        }

        class Query:
            def __init__(self, table):
                self.table_name = table
                self.filters = {}

            def select(self, _columns):
                return self

            def eq(self, column, value):
                self.filters[column] = value
                return self

            def limit(self, _count):
                return self

            def execute(self):
                if self.table_name == "app_access_users":
                    rows = (
                        [account]
                        if self.filters.get("email") == account["email"]
                        else []
                    )
                else:
                    rows = [
                        {"id": app_id, "enabled": True}
                        for app_id in app_ids
                    ]
                return type("Response", (), {"data": rows})()

        class Supabase:
            def table(self, name):
                return Query(name)

        root = Path(__file__).resolve().parents[1]
        app = AppTest.from_file(str(root / "app.py"))
        app.session_state["authenticated_email"] = account["email"]
        with patch("services.db.init_connection", return_value=Supabase()):
            app.run()
        self.assertFalse(app.exception)
        card_titles = {item.value for item in app.subheader}
        self.assertIn(app_ids["qr_generator"], card_titles)
        self.assertIn(app_ids["json_formatter"], card_titles)
        self.assertNotIn(app_ids["pdf_toolbox"], card_titles)


class StreamlitPageTests(unittest.TestCase):
    def setUp(self):
        self.permission_patch = patch(
            "services.access_control.require_module_access",
            return_value=None,
        )
        self.permission_patch.start()

    def tearDown(self):
        self.permission_patch.stop()

    def test_user_management_page_renders_account_forms_without_duplicate_keys(self):
        from streamlit.testing.v1 import AppTest

        class EmptyQuery:
            def select(self, _columns):
                return self

            def order(self, _column):
                return self

            def execute(self):
                return type("Response", (), {"data": []})()

        class EmptySupabase:
            def table(self, _name):
                return EmptyQuery()

        root = Path(__file__).resolve().parents[1]
        with patch("services.db.init_connection", return_value=EmptySupabase()):
            app = AppTest.from_file(
                str(root / "modules/gestion/views/5_usuarios.py")
            )
            app.session_state["authenticated_email"] = "admin@example.com"
            app.run()

        self.assertFalse(app.exception)

    def test_json_page_formats_input_and_reports_invalid_json(self):
        from streamlit.testing.v1 import AppTest

        root = Path(__file__).resolve().parents[1]
        app = AppTest.from_file(
            str(root / "modules/json_formatter/views/app.py")
        ).run()
        self.assertFalse(app.exception)
        app.text_area(key="json_source").set_value('{"value":1}').run()
        app.button(key="json_format").click().run()
        self.assertFalse(app.exception)
        self.assertEqual(app.text_area[1].value, '{\n  "value": 1\n}')
        app.text_area(key="json_source").set_value("{").run()
        app.button(key="json_validate").click().run()
        self.assertFalse(app.exception)
        self.assertTrue(any("línea" in alert.value for alert in app.error))

    def test_qr_and_pdf_pages_render(self):
        from streamlit.testing.v1 import AppTest

        root = Path(__file__).resolve().parents[1]
        qr_page = AppTest.from_file(
            str(root / "modules/qr_generator/views/app.py")
        ).run()
        pdf_page = AppTest.from_file(
            str(root / "modules/pdf_toolbox/views/app.py")
        ).run()
        self.assertFalse(qr_page.exception)
        self.assertFalse(pdf_page.exception)
        self.assertTrue(any("QR Generator" in item.value for item in qr_page.title))
        self.assertTrue(any("PDF Toolbox" in item.value for item in pdf_page.title))
        qr_page.text_area[0].set_value("https://example.com").run()
        qr_page.button(key="generate_qr").click().run()
        self.assertFalse(qr_page.exception)
        self.assertTrue(qr_page.get("download_button"))


if __name__ == "__main__":
    unittest.main()
